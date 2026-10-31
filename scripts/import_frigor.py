"""Prepare FRIGOR's private roster. Source IDs are text, never internal clock numbers.

Reads the supplied workbook without changing it. No personal data belongs in GitHub.
"""
from pathlib import Path
import argparse
import hashlib
import json
import re
import unicodedata
import openpyxl

ROOT = Path(__file__).resolve().parents[1]


def text(value):
    return '' if value is None else str(value).strip()


def identifier(value):
    # Excel numeric identifiers may already have lost leading zeros; do not guess them.
    if not isinstance(value, str) or not re.fullmatch(r'[0-9]{1,30}', value):
        raise ValueError('El carnet debe estar guardado como texto de dígitos, sin espacios.')
    return value


def prepare(rows):
    people, source, seen = [], [], set()
    for row_number, row in rows:
        if not any(v is not None for v in row):
            continue
        if len(row) != 6:
            raise ValueError(f'Fila {row_number}: se esperaban seis columnas.')
        ci = identifier(row[0])
        if ci in seen:
            raise ValueError(f'Fila {row_number}: carnet repetido. No se generó un padrón parcial.')
        seen.add(ci)
        state, permission = text(row[3]), text(row[4])
        if state not in ('Habilitado', 'Deshabilitado') or permission != 'Usuario':
            raise ValueError(f'Fila {row_number}: revisa el estado o permiso antes de importar.')
        original = text(row[1])
        name = ' '.join(original.split())
        missing = not any(c.isalpha() for c in name)
        check_ci = len(ci) not in (7, 8)
        people.append(dict(id='FRIGOR-CI-' + ci, documento=ci, documentoTipo='CI',
                           sap='', nombre=('CI ' + ci + ' · Nombre pendiente') if missing else name,
                           nombreOriginal=original, nombreCompleto='', nombrePendiente=missing,
                           documentoPorRevisar=check_ci, numeroInterno=text(row[2]),
                           planta='FRIGOR', empresa='FRIGOR', centroCosto='', tipo='TRABAJADOR',
                           activo=state == 'Habilitado', qrHash='', vence='',
                           comidas='desayuno,almuerzo,cena'))
        source.append(dict(fila=row_number, carnet=ci, nombre=original, numeroInterno=text(row[2]),
                           estado=state, permiso=permission, grupo=text(row[5]),
                           nombrePendiente=missing, carnetPorRevisar=check_ci))
    if not people:
        raise ValueError('El archivo no contiene usuarios.')
    return people, source


def load(path):
    book = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        if 'Usuarios' not in book.sheetnames:
            raise ValueError('Falta la hoja Usuarios.')
        sheet = book['Usuarios']
        headers = [text(c.value) for c in sheet[4]]
        if len(headers) != 6 or [headers[i] for i in (0, 1, 3, 4, 5)] != ['ID', 'Nombre', 'Estado', 'Permiso', 'Grupo']:
            raise ValueError('No coincide la estructura del archivo USUARIO.xlsx.')
        return prepare((i, tuple(c.value for c in row)) for i, row in enumerate(sheet.iter_rows(min_row=5), 5))
    finally:
        book.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('source', type=Path, nargs='?', default=ROOT / 'USUARIO.xlsx')
    parser.add_argument('--output', type=Path, default=ROOT / 'private/frigor')
    args = parser.parse_args()
    people, source = load(args.source)
    summary = dict(empresa='FRIGOR', origen=args.source.name, hoja='Usuarios', rango='A5:F' + str(source[-1]['fila']),
                   sha256=hashlib.sha256(args.source.read_bytes()).hexdigest(), usuarios=len(people),
                   carnetsDuplicados=0, nombresVacios=sum(not r['nombre'] for r in source),
                   nombresSoloNumeros=sum(bool(r['nombre']) and r['nombrePendiente'] for r in source),
                   nombresDisponibles=sum(not r['nombrePendiente'] for r in source),
                   filasCarnetPorRevisar=[r['fila'] for r in source if r['carnetPorRevisar']],
                   nombresCompletosConfirmados=0, incluyeHuellas=False)
    args.output.mkdir(parents=True, exist_ok=True)
    for name, value in [('padron.json', people), ('origen.json', source), ('revision.json', summary)]:
        target = args.output / name
        temporary = target.with_suffix('.tmp')
        temporary.write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
        temporary.replace(target)
    print(json.dumps(summary, ensure_ascii=True))


if __name__ == '__main__':
    main()
